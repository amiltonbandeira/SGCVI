from datetime import date, timedelta
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Inches, Pt, RGBColor
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


OUT = Path("/home/amilton.bandeira/n360_personalized_saas")
TODAY = date(2026, 10, 3)
PURPLE = "5534D8"
BLUE = "1F4E78"
GREEN = "70AD47"
LIGHT = "EAF2F8"
ORANGE = "F4B183"
RED = "F4CCCC"
WHITE = "FFFFFF"
THIN = Side(style="thin", color="D9E1F2")


def excel_title(ws, title, end_col):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end_col)
    cell = ws.cell(1, 1, title)
    cell.font = Font(size=18, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=PURPLE)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32


def style_header(row):
    for cell in row:
        cell.font = Font(bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(bottom=THIN)


def create_workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Programa diário"
    excel_title(ws, "PROGRAMA DIÁRIO | 30 DIAS", 4)
    ws["A2"] = "Nota de saúde"
    ws["B2"] = "A versão pedida tem 5h30 de sono. Use apenas como emergência; a meta segura é 7–8h."
    ws["A2"].font = Font(bold=True, color="9C0006")
    ws["B2"].font = Font(color="9C0006")
    rows = [
        ["Hora", "Bloco", "O que fazer", "Concluído"],
        ["05:30–06:00", "Arranque", "Água, higiene, cama feita e pequeno-almoço simples", "☐"],
        ["06:00–20:00", "Universidade + estágio", "Deslocação, aulas, estágio e regresso a casa", "☐"],
        ["20:00–21:00", "Chegada", "Banho, jantar e 30 min de descanso sem álcool", "☐"],
        ["21:00–21:30", "Leitura", "Ler 30 minutos", "☐"],
        ["21:30–22:00", "Fé", "Adorar/orar e refletir por 30 minutos", "☐"],
        ["22:00–23:00", "Fundamentos", "10 min CDI • 10 min FIS • 10 min Inglês • 10 min Educação Financeira • 10 min Educação Sexual + 10 min revisão", "☐"],
        ["23:00–00:00", "Programação", "Prática deliberada na área profissional escolhida", "☐"],
        ["00:00–05:30", "Sono (provisório)", "Dormir. Substituir progressivamente por 22:30–05:30 ou outro horário com ≥7h.", "☐"],
    ]
    for row in rows:
        ws.append(row)
    style_header(ws[3])
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN)
        if row[0].row % 2 == 0:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor=LIGHT)
    widths = [16, 22, 95, 14]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.print_area = "A1:D12"

    week = wb.create_sheet("Semana")
    excel_title(week, "SEMANA EM FOCO", 9)
    week.append(["Prioridade", "Seg", "Ter", "Qua", "Qui", "Sex", "Sáb", "Dom", "Resultado da semana"])
    style_header(week[2])
    priorities = [
        ["Sono ≥7h", "", "", "", "", "", "", "", ""],
        ["Sem álcool", "", "", "", "", "", "", "", ""],
        ["Estudo (min)", "", "", "", "", "", "", "", ""],
        ["Programação (min)", "", "", "", "", "", "", "", ""],
        ["Despesas registadas", "", "", "", "", "", "", "", ""],
        ["Exercício/higiene", "", "", "", "", "", "", "", ""],
    ]
    for r in priorities:
        week.append(r)
    for row in week.iter_rows(min_row=3, max_row=week.max_row):
        for c in row:
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = Border(bottom=THIN)
    for col in range(2, 9):
        week.cell(3, col).fill = PatternFill("solid", fgColor=LIGHT)
    for i in range(1, 10):
        week.column_dimensions[get_column_letter(i)].width = 19 if i > 1 else 24
    week.freeze_panes = "B3"

    tracker = wb.create_sheet("30 dias")
    excel_title(tracker, "RASTREADOR DE 30 DIAS", 11)
    tracker.append(["Dia", "Data", "Sem álcool", "Acordei no horário", "Estudei", "Programação", "Gastei com plano", "Sono", "Energia 1–5", "Vitória do dia", "Próxima ação"])
    style_header(tracker[2])
    for i in range(30):
        d = TODAY + timedelta(days=i)
        tracker.append([i + 1, d, "☐", "☐", "☐", "☐", "☐", "", "", "", ""])
    for row in tracker.iter_rows(min_row=3, max_row=32):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN)
        row[1].number_format = "dd/mm/yyyy"
    for width, col in [(8, 1), (14, 2), (14, 3), (18, 4), (14, 5), (16, 6), (18, 7), (12, 8), (14, 9), (28, 10), (30, 11)]:
        tracker.column_dimensions[get_column_letter(col)].width = width
    tracker.freeze_panes = "A3"
    tracker.auto_filter.ref = "A2:K32"

    budget = wb.create_sheet("Dinheiro")
    excel_title(budget, "CONTROLO FINANCEIRO", 6)
    budget.append(["Data", "Categoria", "Descrição", "Valor", "Necessário? (S/N)", "Observação"])
    style_header(budget[2])
    for _ in range(40):
        budget.append(["", "", "", "", "", ""])
    for row in budget.iter_rows(min_row=3, max_row=42):
        for cell in row:
            cell.border = Border(bottom=THIN)
    for col, width in enumerate([14, 22, 38, 15, 20, 35], 1):
        budget.column_dimensions[get_column_letter(col)].width = width
    budget["H2"] = "Resumo"
    budget["H2"].font = Font(bold=True, color=WHITE)
    budget["H2"].fill = PatternFill("solid", fgColor=BLUE)
    budget["H3"] = "Rendimento mensal"
    budget["I3"] = 130000
    budget["H4"] = "Total registado"
    budget["I4"] = "=SUM(D3:D42)"
    budget["H5"] = "Saldo estimado"
    budget["I5"] = "=I3-I4"
    for cell in ["H3", "H4", "H5"]:
        budget[cell].font = Font(bold=True)
    for cell in ["I3", "I4", "I5"]:
        budget[cell].number_format = '#,##0 "Kz"'
    budget.conditional_formatting.add("I5", CellIsRule(operator="lessThan", formula=["0"], fill=PatternFill("solid", fgColor=RED)))

    subjects = wb.create_sheet("Disciplinas")
    excel_title(subjects, "MAPA DE DISCIPLINAS E NOTAS", 9)
    subjects.append(["Disciplina", "Tipo", "Estado atual", "Nota atual", "Meta", "Próxima avaliação", "Próxima tarefa", "Horas/semana", "Concluído"])
    style_header(subjects[2])
    subject_rows = [
        ["ALGLGA", "Universidade", "A confirmar", "", 20, "", "Obter programa e lista de exercícios", 3, "☐"],
        ["CDI I", "Universidade", "A confirmar", "", 20, "", "Obter programa e exercícios", 3, "☐"],
        ["CDI III", "Universidade", "A confirmar", "", 20, "", "Confirmar se é disciplina atual", 2, "☐"],
        ["FISG II", "Universidade", "A confirmar", "", 20, "", "Obter programa e exercícios", 3, "☐"],
        ["EST", "Universidade", "A confirmar", "", 20, "", "Obter programa e avaliações", 2, "☐"],
        ["Inglês", "Fundamento", "A iniciar", "", 20, "", "Diagnóstico de nível", 2, "☐"],
        ["Educação financeira", "Vida", "A iniciar", "", 20, "", "Montar orçamento mensal", 1, "☐"],
        ["Educação sexual", "Vida", "A iniciar", "", 20, "", "Estudo responsável e saúde", 1, "☐"],
        ["Programação/Odoo", "Carreira", "Em desenvolvimento", "", 20, "", "Projeto prático semanal", 5, "☐"],
    ]
    for row in subject_rows:
        subjects.append(row)
    for row in subjects.iter_rows(min_row=3, max_row=subjects.max_row):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN)
    for col, width in enumerate([22, 16, 20, 14, 10, 20, 42, 15, 14], 1):
        subjects.column_dimensions[get_column_letter(col)].width = width
    subjects.freeze_panes = "A3"
    for sheet in wb.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(path)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level)
    p.style.font.color.rgb = RGBColor.from_string(PURPLE)
    return p


def add_bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def create_docx(path):
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.75)
    sec.right_margin = Inches(0.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    for name in ["Title", "Heading 1", "Heading 2"]:
        doc.styles[name].font.name = "Aptos Display"
    title = doc.add_heading("PLANO DE RECONSTRUÇÃO E CRESCIMENTO", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.runs[0].font.color.rgb = RGBColor.from_string(PURPLE)
    p = doc.add_paragraph("Amilton | versão inicial de 30 dias • 03/10/2026")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True

    add_heading(doc, "1. Princípio", 1)
    doc.add_paragraph(
        "O objetivo não é parecer perfeito. É tornar-te estável, confiável e competente através de pequenas ações repetidas. "
        "Este documento é um sistema de execução, não uma promessa de transformação instantânea."
    )
    doc.add_paragraph("Prioridade dos próximos 30 dias: proteger a recuperação, recuperar controlo do tempo, estudar com método e conhecer o dinheiro.")

    add_heading(doc, "2. Regra de saúde e recuperação", 1)
    add_bullets(doc, [
        "30 dias sem álcool: de 03/10/2026 a 01/11/2026.",
        "Evitar ambientes e pessoas que normalmente levam à embriaguez. Se a vontade surgir, sair do local, comer, beber água e contactar uma pessoa segura.",
        "Se surgirem tremores fortes, confusão, convulsões, alucinações ou mal-estar intenso, procurar atendimento médico urgente.",
        "Sono: a rotina pedida de 00:00–05:30 oferece apenas 5h30. A meta de desempenho é pelo menos 7h; ajustar progressivamente para dormir por volta das 22:30.",
    ])

    add_heading(doc, "3. Rotina diária", 1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    for cell, text in zip(table.rows[0].cells, ["Hora", "Bloco", "Execução"]):
        cell.text = text
    daily = [
        ("05:30–06:00", "Arranque", "Água, higiene, cama feita e pequeno-almoço."),
        ("06:00–20:00", "Universidade + estágio", "Deslocação, aulas, estágio e regresso."),
        ("20:00–21:00", "Recuperação", "Banho, jantar e descanso sem álcool."),
        ("21:00–21:30", "Leitura", "30 minutos; registar uma ideia útil."),
        ("21:30–22:00", "Fé", "Adoração, oração e reflexão."),
        ("22:00–23:00", "Fundamentos", "10 min CDI, FIS, Inglês, educação financeira e educação sexual; 10 min de revisão."),
        ("23:00–00:00", "Programação", "Prática num projeto ou competência da área escolhida."),
        ("00:00–05:30", "Sono provisório", "Usar apenas enquanto se corrige o horário; prioridade é chegar a ≥7h."),
    ]
    for values in daily:
        cells = table.add_row().cells
        for cell, text in zip(cells, values):
            cell.text = text
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    add_heading(doc, "4. Metas com prazos", 1)
    metas = [
        ("Até 09/10/2026", "Criar mapa de todas as disciplinas, professores, conteúdos, avaliações e pendências."),
        ("Até 10/10/2026", "Registar todas as despesas dos últimos 7 dias e começar o orçamento no Excel."),
        ("Até 11/10/2026", "Definir uma pessoa de confiança para contactar em caso de vontade de beber."),
        ("Até 17/10/2026", "Cumprir 7 dias sem álcool, estudar pelo menos 5 dias e ter uma rotina de sono observável."),
        ("Até 01/11/2026", "Completar 30 dias sem álcool, fechar o primeiro orçamento e entregar as tarefas urgentes."),
        ("Até 03/12/2026", "Ter uma rotina de estudo sustentável, notas/pendências mapeadas e um projeto pequeno de programação."),
        ("Até 03/04/2027", "Concluir o semestre com todas as avaliações preparadas e um portfólio/projeto demonstrável."),
        ("12–24 meses", "Avançar no curso, aumentar competência e rendimento, formar reserva financeira e melhorar progressivamente a habitação."),
    ]
    mt = doc.add_table(rows=1, cols=2)
    mt.style = "Light Shading Accent 1"
    mt.rows[0].cells[0].text = "Prazo"
    mt.rows[0].cells[1].text = "Meta verificável"
    for deadline, goal in metas:
        cells = mt.add_row().cells
        cells[0].text, cells[1].text = deadline, goal

    add_heading(doc, "5. Plano para notas máximas", 1)
    doc.add_paragraph(
        "Nota máxima é uma meta ambiciosa, mas não se garante apenas com vontade. Primeiro obtém os critérios oficiais de cada disciplina; "
        "o Excel já contém um mapa inicial baseado no horário visível no portal (ALGLGA, CDI I/III, FISG II e EST). Confirma os nomes e acrescenta o que faltar."
    )
    add_bullets(doc, [
        "Para cada disciplina: recolher programa, bibliografia, calendário, pesos das avaliações, provas antigas e lista de exercícios.",
        "Antes de cada aula: 15 minutos para pré-visualizar conceitos e escrever duas perguntas.",
        "Depois de cada aula: 24 horas para criar uma página de resumo e resolver pelo menos dois exercícios.",
        "Sábado: bloco de 2 horas para a disciplina mais fraca; domingo: revisão semanal sem consultar os apontamentos.",
        "A cada 14 dias: fazer uma prova simulada com tempo limitado, corrigir e manter um caderno de erros.",
        "Pedir feedback ao professor/monitor cedo, não apenas na véspera da avaliação.",
    ])
    study = doc.add_table(rows=1, cols=3)
    study.style = "Light Shading Accent 1"
    for cell, text in zip(study.rows[0].cells, ["Área", "Método semanal", "Evidência de domínio"]):
        cell.text = text
    study_rows = [
        ("ALGLGA / CDI", "Definições + exercícios graduais + prova cronometrada.", "Resolver sem olhar e explicar o raciocínio."),
        ("FISG II", "Fórmulas, unidades, exemplos e problemas mistos.", "Identificar dados, modelo, cálculo e unidade final."),
        ("EST", "Conceitos, exemplos e interpretação de resultados.", "Resolver e interpretar, não apenas calcular."),
        ("Inglês", "Vocabulário, leitura curta, escuta e escrita.", "Resumo oral/escrito de um texto."),
        ("Programação/Odoo", "Projeto pequeno, documentação e revisão de código.", "Funcionalidade executável e explicada."),
    ]
    for values in study_rows:
        cells = study.add_row().cells
        for cell, text in zip(cells, values):
            cell.text = text

    add_heading(doc, "6. Dinheiro e estabilidade", 1)
    add_bullets(doc, [
        "Confirmar se os 130.000 são kwanzas e listar despesas fixas: transporte, alimentação, telefone, propinas, casa e dívidas.",
        "Registar cada gasto no Excel no mesmo dia.",
        "Separar primeiro necessidades, depois reserva, e só então lazer.",
        "Começar uma reserva com 5% do rendimento, aumentando quando o orçamento permitir.",
        "Não usar álcool, empréstimos impulsivos ou compras como estratégia de aliviar stress.",
        "Meta de médio prazo: reserva equivalente a pelo menos um mês de despesas essenciais.",
    ])

    add_heading(doc, "7. Revisão semanal de domingo", 1)
    add_bullets(doc, [
        "O que funcionou? O que falhou sem desculpas?",
        "Quantos dias sem álcool? Quantos minutos de estudo real?",
        "Qual é a disciplina mais perigosa esta semana?",
        "Qual é a única tarefa que, se concluída, torna a próxima semana melhor?",
        "Preparar mochila, roupa, alimentação, calendário e orçamento.",
    ])
    doc.add_paragraph("Assinatura semanal: ____________________    Data: ____/____/______")

    add_heading(doc, "8. Nota final", 1)
    doc.add_paragraph(
        "Tu não precisas resolver toda a tua história antes de começares. Precisas de repetir o próximo comportamento correto. "
        "Se houver recaída, interrompe o ciclo rapidamente, procura apoio humano/profissional e volta ao plano sem transformar um erro numa identidade."
    )
    doc.save(path)


if __name__ == "__main__":
    create_workbook(OUT / "Programa_Pessoal_30_Dias.xlsx")
    create_docx(OUT / "Plano_de_Reconstrucao_e_Objectivos.docx")
    print("created")
